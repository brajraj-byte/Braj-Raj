import os
import json
import tempfile

from emergentintegrations.llm.chat import LlmChat, UserMessage, FileContentWithMimeType

EMERGENT_LLM_KEY = os.environ.get("EMERGENT_LLM_KEY")
PARSE_MODEL = ("gemini", "gemini-2.5-flash")

FIELDS = [
    "name", "mobile_number", "email", "age", "date_of_birth", "city", "ctc",
    "industry", "present_company", "designation", "preferred_location",
    "highest_qualification", "linkedin_url", "skill_sets", "notice_period",
]

SYSTEM_MSG = (
    "You are an expert resume/CV parser for a recruitment agency. "
    "Extract structured candidate data accurately. Always respond with ONLY valid JSON, no markdown."
)


def _docx_to_text(path: str) -> str:
    import docx
    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _html_to_text(raw: str) -> str:
    from html.parser import HTMLParser

    class _TextExtractor(HTMLParser):
        def __init__(self):
            super().__init__()
            self.skip = 0
            self.parts = []

        def handle_starttag(self, tag, attrs):
            if tag in ("script", "style"):
                self.skip += 1

        def handle_endtag(self, tag):
            if tag in ("script", "style") and self.skip:
                self.skip -= 1

        def handle_data(self, data):
            if not self.skip and data.strip():
                self.parts.append(data.strip())

    p = _TextExtractor()
    p.feed(raw)
    return "\n".join(p.parts)


def _extract_text(path: str, ext: str, file_bytes: bytes) -> str:
    if ext in ("docx", "doc"):
        try:
            return _docx_to_text(path)
        except Exception:
            return file_bytes.decode("utf-8", "ignore")
    raw = file_bytes.decode("utf-8", "ignore")
    if ext in ("html", "htm"):
        return _html_to_text(raw)
    return raw


def _prompt() -> str:
    return (
        "Parse this resume and return a JSON object with EXACTLY these keys: "
        + ", ".join(FIELDS) + ". "
        "Rules: 'skill_sets' must be an array of strings. "
        "'age' should be a number (integer) if derivable from date of birth or stated age, else null. "
        "'date_of_birth' in DD-MM-YYYY or YYYY-MM-DD if present else null. "
        "'ctc' as a short string like '12 LPA' if present else null. "
        "'notice_period' as a short string like '30 days' if present else null. "
        "Use null for any field not found. Do not invent data. Return ONLY the JSON object."
    )


def _clean_json(text: str) -> dict:
    text = text.strip()
    if text.startswith("```"):
        text = text.split("```", 2)[1]
        if text.lstrip().startswith("json"):
            text = text.lstrip()[4:]
    text = text.strip().strip("`").strip()
    start = text.find("{")
    end = text.rfind("}")
    if start != -1 and end != -1:
        text = text[start:end + 1]
    return json.loads(text)


async def parse_resume(file_bytes: bytes, ext: str) -> dict:
    ext = ext.lower().lstrip(".")
    chat = LlmChat(
        api_key=EMERGENT_LLM_KEY,
        session_id="resume-parse",
        system_message=SYSTEM_MSG,
    ).with_model(*PARSE_MODEL)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}")
    tmp.write(file_bytes)
    tmp.flush()
    tmp.close()
    try:
        if ext == "pdf":
            file_content = FileContentWithMimeType(file_path=tmp.name, mime_type="application/pdf")
            msg = UserMessage(text=_prompt(), file_contents=[file_content])
        else:
            text = _extract_text(tmp.name, ext, file_bytes)
            msg = UserMessage(text=_prompt() + "\n\nRESUME TEXT:\n" + text[:20000])
        resp = await chat.send_message(msg)
    finally:
        try:
            os.unlink(tmp.name)
        except Exception:
            pass

    raw = resp if isinstance(resp, str) else getattr(resp, "content", str(resp))
    try:
        data = _clean_json(raw)
    except Exception:
        data = {}
    result = {}
    for f in FIELDS:
        result[f] = data.get(f)
    if isinstance(result.get("skill_sets"), str):
        result["skill_sets"] = [s.strip() for s in result["skill_sets"].split(",") if s.strip()]
    if not isinstance(result.get("skill_sets"), list):
        result["skill_sets"] = []
    return result


def extract_photo(file_bytes: bytes, ext: str):
    """Best-effort: return (image_bytes, ext) of a candidate photo, or (None, None)."""
    ext = ext.lower().lstrip(".")
    try:
        if ext == "pdf":
            import fitz
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            best = None
            for page in doc:
                for img in page.get_images(full=True):
                    xref = img[0]
                    base = doc.extract_image(xref)
                    data = base["image"]
                    w = base.get("width", 0)
                    h = base.get("height", 0)
                    if w < 60 or h < 60:
                        continue
                    ratio = h / w if w else 0
                    if 0.7 <= ratio <= 2.0 and (best is None or len(data) > best[2]):
                        best = (data, base.get("ext", "png"), len(data))
            if best:
                return best[0], best[1]
        elif ext in ("docx", "doc"):
            import docx
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}")
            tmp.write(file_bytes)
            tmp.close()
            d = docx.Document(tmp.name)
            os.unlink(tmp.name)
            best = None
            for rel in d.part.rels.values():
                if "image" in rel.reltype:
                    blob = rel.target_part.blob
                    if best is None or len(blob) > len(best[0]):
                        ctype = rel.target_part.content_type
                        iext = ctype.split("/")[-1]
                        best = (blob, iext)
            if best:
                return best
    except Exception:
        return None, None
    return None, None
